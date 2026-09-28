"""End-to-end pipeline behaviour, including the rule that defines the product:
no phase runs until a human approves the one before it."""
from __future__ import annotations

import zipfile
from pathlib import Path

import pytest

from core.models import GateStatus, Phase

from .conftest import UC1, UC2


def test_routing_sends_config_work_to_3a_only(pipeline):
    state = pipeline.start(title="UC1", business_process="CU", text=UC1)
    assert state.change_items
    assert Phase.BUILD_INTEGRATION in state.skipped_phases
    assert Phase.BUILD_CONFIG not in state.skipped_phases
    assert all(i.build_owner == "3A" for i in state.change_items)


def test_routing_sends_integration_work_to_3b(pipeline):
    state = pipeline.start(title="UC2", business_process="CU", text=UC2)
    owners = {i.build_owner for i in state.change_items}
    assert "3B" in owners
    assert Phase.BUILD_INTEGRATION not in state.skipped_phases


def test_gate_blocks_the_next_phase(pipeline):
    state = pipeline.start(title="Gate test", business_process="CU", text=UC1)

    first = pipeline.advance(state.run_id)
    assert first.ran and first.result.ok
    assert first.phase is Phase.FDD

    blocked = pipeline.advance(state.run_id)
    assert not blocked.ran, "a phase must not run while the previous gate is open"
    assert blocked.blocked_by == Phase.FDD.value

    pipeline.approve(state.run_id, Phase.FDD, by="tester")
    after = pipeline.advance(state.run_id)
    assert after.ran and after.phase is Phase.TDD


def test_revision_reruns_the_same_phase(pipeline):
    state = pipeline.start(title="Revision test", business_process="CU", text=UC1)
    pipeline.advance(state.run_id)
    pipeline.revise(state.run_id, Phase.FDD, by="tester", comment="add security section")

    again = pipeline.advance(state.run_id)
    assert again.phase is Phase.FDD, "a revision must re-run that phase, not move on"


def test_approval_writes_into_the_brain(pipeline):
    state = pipeline.start(title="Brain write test", business_process="CU", text=UC1)
    before = pipeline.brain.store.stats()["documents"]
    pipeline.advance(state.run_id)
    pipeline.approve(state.run_id, Phase.FDD, by="tester")
    after = pipeline.brain.store.stats()["documents"]
    assert after == before + 1, "an approved artifact must land in the AI Brain"


def test_unapproved_work_is_not_written_to_the_brain(pipeline):
    state = pipeline.start(title="No write test", business_process="CU", text=UC1)
    before = pipeline.brain.store.stats()["documents"]
    pipeline.advance(state.run_id)
    assert pipeline.brain.store.stats()["documents"] == before, (
        "nothing may reach the AI Brain before the gate is approved"
    )


@pytest.mark.parametrize("payload,label", [(UC1, "uc1"), (UC2, "uc2")])
def test_full_run_produces_a_valid_package(pipeline, payload, label):
    state = pipeline.start(title=f"Full {label}", business_process="CU", text=payload)
    steps = pipeline.run_all(state.run_id, auto_approve=True)

    assert all(s.result.ok for s in steps if s.ran), [
        s.result.error for s in steps if s.ran and not s.result.ok
    ]
    status = pipeline.status(state.run_id)
    assert status["finished"], "the run should complete when every gate is auto-approved"
    assert status["total_artifacts"] > 8

    zips = [a for a in pipeline.store.load(state.run_id).artifacts() if a.kind == "zip"]
    assert zips, "Agent 5 must produce a migration package"
    with zipfile.ZipFile(zips[0].path) as zf:
        assert zf.testzip() is None
        names = zf.namelist()
        assert "MANIFEST.json" in names
        assert "RUNBOOK.md" in names


def test_generated_documents_open_cleanly(pipeline):
    state = pipeline.start(title="Artifact validity", business_process="CU", text=UC1)
    pipeline.run_all(state.run_id, auto_approve=True)

    import docx
    import openpyxl

    for artifact in pipeline.store.load(state.run_id).artifacts():
        path = Path(artifact.path)
        assert path.exists(), f"{artifact.name} was registered but not written"
        if artifact.kind == "docx":
            assert len(docx.Document(str(path)).paragraphs) > 3
        elif artifact.kind == "xlsx":
            openpyxl.load_workbook(str(path)).close()
        elif artifact.kind == "xml":
            from xml.etree import ElementTree as ET

            ET.parse(str(path))


def test_integration_run_consolidates_into_one_interface(pipeline):
    """Several requirement lines describing one interface must not become many."""
    import json

    state = pipeline.start(title="Interface consolidation", business_process="CU", text=UC2)
    pipeline.run_all(state.run_id, auto_approve=True)

    spec = next(
        a for a in pipeline.store.load(state.run_id).artifacts() if a.name == "mif_components.json"
    )
    interfaces = json.loads(Path(spec.path).read_text(encoding="utf-8"))
    assert len(interfaces) == 1, f"expected one outbound interface, got {len(interfaces)}"
    assert interfaces[0]["direction"] == "Outbound"
    assert interfaces[0]["publish_channel"]


def test_integration_run_generates_no_jython(pipeline):
    """MIF interfaces are configuration, not automation scripts."""
    state = pipeline.start(title="No scripts for integration", business_process="CU", text=UC2)
    pipeline.run_all(state.run_id, auto_approve=True)
    scripts = [a for a in pipeline.store.load(state.run_id).artifacts() if a.kind == "py"]
    assert not scripts, f"integration work must not emit Jython: {[s.name for s in scripts]}"


def test_estimate_is_per_deliverable_not_per_requirement(pipeline):
    """Six requirement lines about one interface must not cost six interfaces."""
    state = pipeline.start(title="Estimation", business_process="CU", text=UC2)
    pipeline.advance(state.run_id)
    pipeline.approve(state.run_id, Phase.FDD, by="t")
    pipeline.advance(state.run_id)

    items = pipeline.store.load(state.run_id).change_items
    total = sum(i.effort_hours for i in items)
    integration = [i for i in items if i.change_type.value == "integration"]
    assert len(integration) > 1, "this fixture should produce several integration lines"
    assert total < 20 * len(integration), (
        f"{total}h charges close to a full baseline per requirement line"
    )


def test_deploy_excludes_unapproved_phases(pipeline):
    state = pipeline.start(title="Selective deploy", business_process="CU", text=UC1)
    for phase in (Phase.FDD, Phase.TDD, Phase.BUILD_CONFIG):
        pipeline.advance(state.run_id)
        pipeline.approve(state.run_id, phase, by="t")
    pipeline.advance(state.run_id)          # Agent 4 runs
    pipeline.skip(state.run_id, Phase.TEST, by="t", comment="testing deferred")
    result = pipeline.advance(state.run_id)  # Agent 5

    assert result.phase is Phase.DEPLOY
    loaded = pipeline.store.load(state.run_id)
    assert loaded.gate(Phase.TEST).status is GateStatus.SKIPPED
    runbook = next(a for a in result.result.artifacts if a.name == "Deployment_Runbook.md")
    assert Path(runbook.path).exists()
