"""Ingest and rendering: every supported input reads, every output opens."""
from __future__ import annotations

from pathlib import Path

import pytest

from agents.orchestrator import classify_text, extract_requirements
from core.models import ChangeType
from ingest.extract import extract_file, extract_text
from rendering.docx_writer import render
from rendering.package import PackageEntry, build_package
from rendering.xlsx_writer import write_testcases

SAMPLE_MD = """# Requirements
- The system shall add a CU Comment field to the CU Header application.
- An automation script must copy the value to the work order.
"""


def test_text_ingest_detects_user_stories():
    doc = extract_text("As a CU designer I want a comment field.\nAcceptance criteria: visible on CU Header.")
    assert doc.ok
    assert doc.kind == "user_story"


def test_markdown_ingest(tmp_path):
    p = tmp_path / "req.md"
    p.write_text(SAMPLE_MD, encoding="utf-8")
    doc = extract_file(p)
    assert doc.ok and "CU Comment" in doc.text


def test_unsupported_type_fails_softly(tmp_path):
    p = tmp_path / "thing.bin"
    p.write_bytes(b"\x00\x01")
    doc = extract_file(p)
    assert not doc.ok
    assert "Unsupported" in doc.error, "an odd attachment must not crash the run"


def test_missing_file_fails_softly(tmp_path):
    doc = extract_file(tmp_path / "nope.docx")
    assert not doc.ok and "not found" in doc.error.lower()


def test_xlsx_roundtrip(tmp_path):
    import openpyxl

    rows = [
        {
            "Test Case ID": "TC-001", "Module": "CU", "Scenario": "Happy path", "Type": "Happy path",
            "Pre-conditions": "record exists", "Test Steps": "1. open\n2. save",
            "Expected Result": "saved", "Actual Result": "", "Pass/Fail": "", "Notes": "",
        }
    ]
    out = write_testcases(rows, tmp_path / "tc.xlsx", summary_rows=[["CI-1", 1, 0, 0, 0, 1]])
    wb = openpyxl.load_workbook(str(out))
    assert wb.sheetnames == ["Test Cases", "Coverage Summary"]
    assert wb["Test Cases"].max_row == 2
    wb.close()


def test_docx_render_uses_the_template(cfg, tmp_path):
    import docx

    md = "# Title\n\n## Section\ntext with **bold**\n\n| A | B |\n|---|---|\n| 1 | 2 |\n\n- bullet\n"
    result = render(md, template=cfg.template_path("fdd"), target=tmp_path / "out.docx", title="T")
    assert result.tables == 1
    d = docx.Document(str(tmp_path / "out.docx"))
    assert len(d.paragraphs) > 3
    assert len(d.tables) == 1


def test_package_is_reproducible(tmp_path):
    entries = [PackageEntry(arcname="a/x.txt", text="hello"), PackageEntry(arcname="b/y.txt", text="world")]
    first = build_package(entries, tmp_path / "p1.zip", manifest={"run_id": "R1"})
    second = build_package(entries, tmp_path / "p2.zip", manifest={"run_id": "R1"})
    assert first.entries == second.entries
    # Same content -> same checksums, regardless of when it was built.
    assert [f["sha256"] for f in first.manifest["files"]] == [f["sha256"] for f in second.manifest["files"]]


@pytest.mark.parametrize(
    "text,expected",
    [
        ("Add a CU Comment field to the CU Header application", ChangeType.CONFIG),
        ("An automation script shall copy CUJP.CUCOMMENT to WORKORDER.DESCRIPTION", ChangeType.CUSTOMISATION),
        ("Trigger an outbound message to the AUD external system", ChangeType.INTEGRATION),
        ("Define a Publish Channel on the CUJP object structure", ChangeType.INTEGRATION),
        ("Route the record for approval through a workflow process", ChangeType.WORKFLOW),
        ("Produce a BIRT report of CU estimates by substation", ChangeType.REPORT),
        ("Empty CUCOMMENT must not overwrite an existing description", ChangeType.CUSTOMISATION),
    ],
)
def test_classification(text, expected):
    got, confidence = classify_text(text)
    assert got is expected, f"{text!r} -> {got.value}, expected {expected.value}"
    assert confidence > 0.5


def test_requirement_extraction_finds_bullets():
    doc = extract_text(SAMPLE_MD, "req.md")
    reqs = extract_requirements([doc])
    assert len(reqs) == 2
    assert all(r.title for r in reqs)


def test_narrative_prose_is_not_a_change_item():
    """Explanatory background must not become something to build and estimate."""
    doc = extract_text(
        "# Use Case\n\n"
        "## Background\n"
        "The AUD asset utilisation system needs to be notified whenever a CU estimate is accepted.\n"
        "Today planners must be told by email, which delays the work.\n\n"
        "## Requirements\n"
        "- The system shall trigger a CU outbound message to the AUD external system on ACCEPTED.\n",
        "uc.md",
    )
    reqs = extract_requirements([doc])
    assert len(reqs) == 1, [r.text for r in reqs]
    assert "shall trigger" in reqs[0].text
    assert all("Background" not in r.source for r in reqs)
