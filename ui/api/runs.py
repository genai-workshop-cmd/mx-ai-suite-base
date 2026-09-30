"""Run lifecycle endpoints: create, advance, gate decisions, artifacts."""
from __future__ import annotations

import shutil
import tempfile
import threading
from pathlib import Path
from typing import Any

from fastapi import APIRouter, Body, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse, PlainTextResponse

from core.errors import SuiteError
from core.models import Phase
from ingest.extract import SUPPORTED

from .deps import get_config, get_pipeline, get_store, job_create, job_done, job_fail, job_get

router = APIRouter(prefix="/api/runs", tags=["runs"])


def _fail(exc: SuiteError) -> HTTPException:
    return HTTPException(status_code=400, detail=exc.as_dict())


def _phase(name: str) -> Phase:
    try:
        return Phase(name.lower())
    except ValueError:
        raise HTTPException(status_code=400, detail={"code": "PHASE", "message": f"Unknown phase '{name}'."})


@router.get("")
def list_runs(limit: int = 50) -> list[dict]:
    return get_store().list(limit)


@router.post("")
async def create_run(
    title: str = Form(...),
    business_process: str = Form(""),
    text: str = Form(""),
    notes: str = Form(""),
    files: list[UploadFile] = File(default=[]),
) -> dict:
    """Create a run from uploaded files and/or pasted text."""
    staged: list[Path] = []
    tmp = Path(tempfile.mkdtemp(prefix="mxsuite_upload_"))
    try:
        for upload in files or []:
            if not upload.filename:
                continue
            suffix = Path(upload.filename).suffix.lower()
            if suffix not in SUPPORTED:
                raise HTTPException(
                    status_code=400,
                    detail={
                        "code": "INGEST",
                        "message": f"Unsupported file type '{suffix}' ({upload.filename}).",
                        "remedy": f"Supported: {', '.join(sorted(SUPPORTED))}",
                    },
                )
            dest = tmp / Path(upload.filename).name
            dest.write_bytes(await upload.read())
            staged.append(dest)

        if not staged and not text.strip():
            raise HTTPException(
                status_code=400,
                detail={
                    "code": "INGEST",
                    "message": "Provide at least one file or some requirement text.",
                    "remedy": "Attach a document or paste the requirements.",
                },
            )
        try:
            state = get_pipeline().start(
                title=title, business_process=business_process, files=staged, text=text, notes=notes
            )
        except SuiteError as exc:
            raise _fail(exc)
        return get_pipeline().status(state.run_id)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


@router.get("/{run_id}")
def run_status(run_id: str) -> dict:
    try:
        return get_pipeline().status(run_id)
    except SuiteError as exc:
        raise _fail(exc)


@router.post("/{run_id}/advance")
def advance(run_id: str, payload: dict[str, Any] = Body(default={})) -> dict:
    """Start the next phase in a background thread and return a job ID to poll.

    The caller polls GET /{run_id}/job/{job_id} until status is 'done' or 'error'.
    This prevents long-running LLM calls from blocking the HTTP connection.
    """
    job_id = job_create()

    def _worker() -> None:
        pipe = get_pipeline()
        try:
            if payload.get("all"):
                steps = pipe.run_all(run_id, auto_approve=bool(payload.get("auto")))
                result = {"steps": [s.as_dict() for s in steps], "status": pipe.status(run_id)}
            else:
                step = pipe.advance(run_id, auto_approve=bool(payload.get("auto")))
                result = {"steps": [step.as_dict()], "status": pipe.status(run_id)}
            job_done(job_id, result)
        except SuiteError as exc:
            job_fail(job_id, exc.as_dict())
        except Exception as exc:
            job_fail(job_id, {"code": "UNEXPECTED", "message": str(exc), "remedy": "See the run log."})

    threading.Thread(target=_worker, daemon=True, name=f"advance-{run_id[:8]}").start()
    return {"job_id": job_id, "status": "running"}


@router.get("/{run_id}/job/{job_id}")
def poll_job(run_id: str, job_id: str) -> dict:
    """Poll the status of a background advance job."""
    job = job_get(job_id)
    if job is None:
        raise HTTPException(
            status_code=404,
            detail={"code": "JOB", "message": f"Job '{job_id}' not found or already expired."},
        )
    return job


@router.post("/{run_id}/gate/{phase}")
def decide_gate(run_id: str, phase: str, payload: dict[str, Any] = Body(default={})) -> dict:
    """Record a gate decision: approve, revise or skip."""
    pipe = get_pipeline()
    target = _phase(phase)
    decision = str(payload.get("decision", "approve")).lower()
    by = str(payload.get("by", "")).strip()
    comment = str(payload.get("comment", "")).strip()

    try:
        if decision == "approve":
            pipe.approve(run_id, target, by=by, comment=comment,
                         duplicate_action=str(payload.get("duplicate_action", "")))
        elif decision in ("revise", "revision"):
            if not comment:
                raise HTTPException(
                    status_code=400,
                    detail={"code": "GATE", "message": "A revision needs a comment saying what to change."},
                )
            pipe.revise(run_id, target, by=by, comment=comment)
        elif decision == "skip":
            pipe.skip(run_id, target, by=by, comment=comment)
        else:
            raise HTTPException(
                status_code=400,
                detail={"code": "GATE", "message": f"Unknown decision '{decision}'. Use approve, revise or skip."},
            )
    except SuiteError as exc:
        raise _fail(exc)
    return pipe.status(run_id)


@router.get("/{run_id}/artifact")
def artifact(run_id: str, path: str, download: bool = False):
    """Serve one artifact. Paths are confined to the run directory."""
    store = get_store()
    run_dir = store.run_dir(run_id).resolve()
    if not run_dir.exists():
        raise HTTPException(status_code=404, detail={"code": "RUN", "message": f"Run '{run_id}' not found."})

    target = Path(path)
    resolved = (target if target.is_absolute() else run_dir / target).resolve()
    # Path traversal guard: the artifact must live inside this run.
    if run_dir != resolved and run_dir not in resolved.parents:
        raise HTTPException(status_code=403, detail={"code": "PATH", "message": "Artifact is outside the run."})
    if not resolved.is_file():
        raise HTTPException(status_code=404, detail={"code": "PATH", "message": f"Not found: {resolved.name}"})

    if download or resolved.suffix.lower() in {".docx", ".xlsx", ".zip"}:
        return FileResponse(resolved, filename=resolved.name)
    if resolved.suffix.lower() in {".md", ".txt", ".json", ".xml", ".py", ".log"}:
        return PlainTextResponse(resolved.read_text(encoding="utf-8", errors="replace"))
    return FileResponse(resolved, filename=resolved.name)


@router.get("/{run_id}/log")
def run_log(run_id: str, lines: int = 200) -> dict:
    path = get_store().log_path(run_id)
    if not path.exists():
        return {"lines": []}
    tail = path.read_text(encoding="utf-8", errors="replace").splitlines()[-lines:]
    return {"lines": tail}


@router.get("/{run_id}/preview")
def preview(run_id: str, path: str) -> dict:
    """Markdown/text preview for the review panel."""
    response = artifact(run_id, path, download=False)
    if isinstance(response, PlainTextResponse):
        return {"kind": "text", "content": response.body.decode("utf-8", errors="replace")}

    resolved = Path(path)
    if resolved.suffix.lower() == ".docx":
        try:
            import mammoth

            with open(resolved, "rb") as fh:
                return {"kind": "html", "content": mammoth.convert_to_html(fh).value}
        except Exception as exc:
            return {"kind": "text", "content": f"(preview unavailable: {exc}) - use Download."}
    if resolved.suffix.lower() == ".xlsx":
        return {"kind": "table", "content": _xlsx_preview(resolved)}
    return {"kind": "binary", "content": ""}


def _xlsx_preview(path: Path, max_rows: int = 60) -> dict:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    try:
        sheets = []
        for ws in wb.worksheets:
            rows = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= max_rows:
                    break
                rows.append(["" if v is None else str(v) for v in row])
            sheets.append({"name": ws.title, "rows": rows, "truncated": ws.max_row > max_rows})
        return {"sheets": sheets}
    finally:
        wb.close()


@router.delete("/{run_id}")
def delete_run(run_id: str) -> dict:
    get_store().delete(run_id)
    return {"deleted": run_id}


@router.get("/meta/processes")
def processes() -> dict:
    cfg = get_config()
    out = []
    for name in cfg.known_processes():
        try:
            p = cfg.process(name)
        except SuiteError:
            continue
        out.append({"name": name, "label": p.get("label", name), "description": p.get("description", "").strip()})
    return {"processes": out, "default": cfg.default_process, "supported_files": sorted(SUPPORTED)}
