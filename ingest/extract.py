"""Turn whatever the user uploads into normalised text the agents can read.

Blueprint section 4, Agent 1 inputs: ADO/user stories, business decisions,
whiteboard notes, Teams transcripts, meeting notes, "or any combination".

Every extractor is defensive: an unreadable file yields a Document with an
`error` set rather than aborting the run, so one bad attachment cannot stop
a pipeline.
"""
from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from core.errors import IngestError
from core.logging import get

log = get("suite.ingest")

SUPPORTED = {".docx", ".xlsx", ".xlsm", ".pdf", ".pptx", ".md", ".txt", ".json", ".csv", ".html", ".htm"}

#: Heuristics that tell us what kind of material this is, which the
#: orchestrator uses when weighting requirements.
_KIND_PATTERNS: list[tuple[str, re.Pattern]] = [
    ("user_story", re.compile(r"\b(as an?\s+\w+.{0,40}\bi want\b|acceptance criteria|user story|story points)\b", re.I)),
    ("ado_export", re.compile(r"\b(work item (id|type)|azure devops|\bADO\b|backlog item)\b", re.I)),
    ("meeting_notes", re.compile(r"\b(attendees?|action items?|minutes of meeting|\bMoM\b|agenda)\b", re.I)),
    ("teams_chat", re.compile(r"^\s*\[?\d{1,2}[:/]\d{2}.{0,20}\]?\s*\w+\s*:", re.M)),
    ("decision_log", re.compile(r"\b(decision (log|register)|agreed that|it was decided|KBD)\b", re.I)),
    ("spec", re.compile(r"\b(functional design|technical design|requirement(s)? specification|FDD|TDD)\b", re.I)),
]


@dataclass
class Document:
    """One ingested input."""

    name: str
    path: str
    kind: str = "unknown"
    text: str = ""
    tables: list[list[list[str]]] = field(default_factory=list)
    metadata: dict[str, Any] = field(default_factory=dict)
    error: str = ""

    @property
    def ok(self) -> bool:
        return not self.error and bool(self.text.strip() or self.tables)

    @property
    def chars(self) -> int:
        return len(self.text)

    def as_dict(self) -> dict[str, Any]:
        return {
            "name": self.name,
            "path": self.path,
            "kind": self.kind,
            "chars": self.chars,
            "tables": len(self.tables),
            "error": self.error,
        }

    def to_prompt_block(self, max_chars: int = 12000) -> str:
        """Render for inclusion in an agent prompt."""
        parts = [f"### SOURCE: {self.name}  (detected type: {self.kind})"]
        body = self.text.strip()
        if len(body) > max_chars:
            body = body[:max_chars] + f"\n...[truncated, {self.chars - max_chars} more characters]"
        parts.append(body)
        for i, table in enumerate(self.tables[:6], 1):
            parts.append(f"\n#### Table {i} from {self.name}")
            parts.append(render_table(table))
        return "\n".join(parts)


def render_table(rows: list[list[str]], max_rows: int = 40) -> str:
    if not rows:
        return ""
    trimmed = rows[:max_rows]
    width = max(len(r) for r in trimmed)
    norm = [list(r) + [""] * (width - len(r)) for r in trimmed]
    out = ["| " + " | ".join(str(c).replace("|", "\\|") for c in norm[0]) + " |"]
    out.append("|" + "---|" * width)
    for row in norm[1:]:
        out.append("| " + " | ".join(str(c).replace("|", "\\|") for c in row) + " |")
    if len(rows) > max_rows:
        out.append(f"_...{len(rows) - max_rows} more rows_")
    return "\n".join(out)


def detect_kind(text: str, suffix: str) -> str:
    for kind, pattern in _KIND_PATTERNS:
        if pattern.search(text):
            return kind
    return {".xlsx": "spreadsheet", ".xlsm": "spreadsheet", ".csv": "spreadsheet",
            ".pptx": "slides", ".pdf": "pdf", ".json": "data"}.get(suffix, "notes")


# --------------------------------------------------------------------------
# extractors
# --------------------------------------------------------------------------
def _docx(path: Path) -> tuple[str, list[list[list[str]]]]:
    import docx

    doc = docx.Document(str(path))
    lines: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower()
        if style.startswith("heading"):
            level = "".join(ch for ch in style if ch.isdigit()) or "1"
            lines.append(f"{'#' * min(int(level), 6)} {text}")
        elif style.startswith("list"):
            lines.append(f"- {text}")
        else:
            lines.append(text)

    tables: list[list[list[str]]] = []
    for table in doc.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if any(any(c for c in r) for r in rows):
            tables.append(rows)
    return "\n\n".join(lines), tables


def _xlsx(path: Path) -> tuple[str, list[list[list[str]]]]:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    tables: list[list[list[str]]] = []
    lines: list[str] = []
    try:
        for sheet in wb.worksheets:
            rows: list[list[str]] = []
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if v is None else str(v).strip() for v in row]
                if any(cells):
                    rows.append(cells)
            if rows:
                lines.append(f"## Sheet: {sheet.title} ({len(rows)} rows)")
                tables.append(rows)
    finally:
        wb.close()
    return "\n\n".join(lines), tables


def _pdf(path: Path) -> tuple[str, list[list[list[str]]]]:
    try:
        from pypdf import PdfReader
    except ImportError:
        try:
            from PyPDF2 import PdfReader  # type: ignore
        except ImportError as exc:
            raise IngestError(
                "Reading PDFs needs pypdf.",
                remedy="pip install pypdf",
            ) from exc

    reader = PdfReader(str(path))
    pages = []
    for i, page in enumerate(reader.pages, 1):
        try:
            text = page.extract_text() or ""
        except Exception as exc:
            log.debug("pdf page %d failed: %s", i, exc)
            continue
        if text.strip():
            pages.append(f"## Page {i}\n{text.strip()}")
    return "\n\n".join(pages), []


def _pptx(path: Path) -> tuple[str, list[list[list[str]]]]:
    """Slide text without a python-pptx dependency - read the XML directly."""
    text_re = re.compile(r"<a:t>(.*?)</a:t>", re.DOTALL)
    slides: list[str] = []
    with zipfile.ZipFile(path) as zf:
        names = sorted(
            (n for n in zf.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)),
            key=lambda n: int(re.findall(r"\d+", n)[-1]),
        )
        for i, name in enumerate(names, 1):
            xml = zf.read(name).decode("utf-8", errors="ignore")
            fragments = [_unescape(t).strip() for t in text_re.findall(xml)]
            body = "\n".join(f for f in fragments if f)
            if body:
                slides.append(f"## Slide {i}\n{body}")
    return "\n\n".join(slides), []


def _html(path: Path) -> tuple[str, list[list[list[str]]]]:
    raw = path.read_text(encoding="utf-8", errors="ignore")
    raw = re.sub(r"<(script|style)[^>]*>.*?</\1>", " ", raw, flags=re.DOTALL | re.I)
    raw = re.sub(r"<br\s*/?>|</p>|</div>|</li>|</tr>", "\n", raw, flags=re.I)
    text = re.sub(r"<[^>]+>", " ", raw)
    text = _unescape(text)
    return re.sub(r"[ \t]{2,}", " ", re.sub(r"\n{3,}", "\n\n", text)).strip(), []


def _csv(path: Path) -> tuple[str, list[list[list[str]]]]:
    raw = path.read_text(encoding="utf-8", errors="ignore")
    try:
        dialect = csv.Sniffer().sniff(raw[:4096])
    except csv.Error:
        dialect = csv.excel
    rows = [r for r in csv.reader(io.StringIO(raw), dialect) if any(c.strip() for c in r)]
    return f"## {path.name} ({len(rows)} rows)", [rows] if rows else []


def _json(path: Path) -> tuple[str, list[list[list[str]]]]:
    raw = path.read_text(encoding="utf-8", errors="ignore")
    try:
        data = json.loads(raw)
    except json.JSONDecodeError:
        return raw[:20000], []
    return json.dumps(data, indent=2, ensure_ascii=False)[:40000], []


def _plain(path: Path) -> tuple[str, list[list[list[str]]]]:
    return path.read_text(encoding="utf-8", errors="ignore"), []


_EXTRACTORS = {
    ".docx": _docx,
    ".xlsx": _xlsx,
    ".xlsm": _xlsx,
    ".pdf": _pdf,
    ".pptx": _pptx,
    ".html": _html,
    ".htm": _html,
    ".csv": _csv,
    ".json": _json,
    ".md": _plain,
    ".txt": _plain,
}

_ENTITIES = {"&amp;": "&", "&lt;": "<", "&gt;": ">", "&quot;": '"', "&#39;": "'", "&nbsp;": " ", "&apos;": "'"}


def _unescape(text: str) -> str:
    for k, v in _ENTITIES.items():
        text = text.replace(k, v)
    return re.sub(r"&#(\d+);", lambda m: chr(int(m.group(1))), text)


# --------------------------------------------------------------------------
# public API
# --------------------------------------------------------------------------
def extract_file(path: str | Path) -> Document:
    path = Path(path)
    doc = Document(name=path.name, path=str(path))
    if not path.exists():
        doc.error = "File not found."
        return doc

    suffix = path.suffix.lower()
    extractor = _EXTRACTORS.get(suffix)
    if extractor is None:
        doc.error = f"Unsupported file type '{suffix}'. Supported: {', '.join(sorted(SUPPORTED))}"
        return doc

    try:
        text, tables = extractor(path)
    except IngestError as exc:
        doc.error = exc.message
        return doc
    except Exception as exc:
        doc.error = f"Could not read {path.name}: {exc}"
        log.warning("ingest failed for %s: %s", path.name, exc)
        return doc

    doc.text = text.strip()
    doc.tables = tables
    doc.kind = detect_kind(doc.text, suffix)
    doc.metadata = {"bytes": path.stat().st_size, "suffix": suffix}
    log.info("ingested %s (%s, %d chars, %d tables)", path.name, doc.kind, doc.chars, len(tables))
    return doc


def extract_text(text: str, name: str = "pasted-text") -> Document:
    """Wrap pasted text as a Document (the UI 'paste requirements' path)."""
    doc = Document(name=name, path="(pasted)", text=text.strip())
    doc.kind = detect_kind(doc.text, ".txt")
    return doc


def extract_all(paths: list[str | Path]) -> list[Document]:
    return [extract_file(p) for p in paths]


def bundle(docs: list[Document], max_chars_each: int = 12000) -> str:
    """Concatenate ingested documents into one prompt-ready block."""
    usable = [d for d in docs if d.ok]
    if not usable:
        return "(no readable input material)"
    return "\n\n---\n\n".join(d.to_prompt_block(max_chars_each) for d in usable)
