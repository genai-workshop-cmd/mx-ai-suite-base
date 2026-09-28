"""Markdown -> production-ready .docx, using the client template.

The template supplies the styles, page setup, headers and footers. We clone it,
strip its body, and write our content using the template's own named styles, so
output matches whatever the client's template looks like. Nothing about the
document structure is hardcoded (blueprint principle: client-agnostic templates).

Supported markdown: headings, paragraphs, bullet/numbered lists, GFM tables,
fenced code blocks, bold/italic/inline-code spans, and horizontal rules.
"""
from __future__ import annotations

import re
import shutil
from dataclasses import dataclass
from pathlib import Path

from core.errors import RenderError
from core.logging import get

log = get("suite.render.docx")

_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
_BULLET_RE = re.compile(r"^\s*[-*+]\s+(.*)$")
_NUMBER_RE = re.compile(r"^\s*\d+[.)]\s+(.*)$")
_TABLE_SEP_RE = re.compile(r"^\s*\|?[\s:\-|]+\|[\s:\-|]*$")
_FENCE_RE = re.compile(r"^\s*```(\w*)\s*$")
_RULE_RE = re.compile(r"^\s*([-*_])\1{2,}\s*$")
_INLINE_RE = re.compile(r"(\*\*.+?\*\*|__.+?__|\*[^*]+?\*|`[^`]+?`)")


@dataclass
class DocxResult:
    path: Path
    paragraphs: int
    tables: int
    template: str

    @property
    def bytes(self) -> int:
        return self.path.stat().st_size if self.path.exists() else 0


def _blank_from_template(template: Path, target: Path):
    """Copy the template then clear its body, keeping styles and section setup."""
    import docx

    target.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(template, target)
    doc = docx.Document(str(target))
    body = doc.element.body
    # Keep the trailing sectPr (page size, margins, headers/footers).
    for child in list(body):
        if child.tag.endswith("}sectPr"):
            continue
        body.remove(child)
    return doc


def _style_or_none(doc, *names: str) -> str | None:
    available = {s.name for s in doc.styles}
    for name in names:
        if name in available:
            return name
    return None


def _add_runs(paragraph, text: str) -> None:
    """Write inline markdown spans as formatted runs."""
    for part in _INLINE_RE.split(text):
        if not part:
            continue
        if (part.startswith("**") and part.endswith("**")) or (part.startswith("__") and part.endswith("__")):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
        else:
            paragraph.add_run(part)


def _split_row(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip().replace("\\|", "|") for c in line.split("|")]


def render(
    markdown: str,
    *,
    template: Path,
    target: Path,
    title: str | None = None,
    subtitle: str | None = None,
) -> DocxResult:
    """Render markdown into a .docx built from `template`."""
    if not template.exists():
        raise RenderError(f"Template not found: {template}")

    try:
        doc = _blank_from_template(template, target)
    except Exception as exc:
        raise RenderError(f"Could not open template {template.name}: {exc}") from exc

    title_style = _style_or_none(doc, "Title", "Heading 1")
    subtitle_style = _style_or_none(doc, "Subtitle", "Intense Quote")
    quote_style = _style_or_none(doc, "Quote", "Intense Quote")
    code_style = _style_or_none(doc, "HTML Preformatted", "No Spacing")
    bullet_style = _style_or_none(doc, "List Bullet", "List Paragraph")
    number_style = _style_or_none(doc, "List Number", "List Paragraph")
    table_style = _style_or_none(
        doc, "Table Grid", "Light Grid Accent 1", "Light List Accent 1", "Medium Shading 1 Accent 1"
    )

    n_para = n_table = 0

    if title:
        p = doc.add_paragraph(style=title_style) if title_style else doc.add_paragraph()
        _add_runs(p, title)
        n_para += 1
    if subtitle:
        p = doc.add_paragraph(style=subtitle_style) if subtitle_style else doc.add_paragraph()
        _add_runs(p, subtitle)
        n_para += 1

    lines = markdown.replace("\r\n", "\n").split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        fence = _FENCE_RE.match(line)
        if fence:
            block: list[str] = []
            i += 1
            while i < len(lines) and not _FENCE_RE.match(lines[i]):
                block.append(lines[i])
                i += 1
            i += 1  # closing fence
            para = doc.add_paragraph(style=code_style) if code_style else doc.add_paragraph()
            run = para.add_run("\n".join(block))
            run.font.name = "Consolas"
            run.font.size = _pt(9)
            n_para += 1
            continue

        if _RULE_RE.match(stripped):
            doc.add_paragraph("_" * 60)
            n_para += 1
            i += 1
            continue

        heading = _HEADING_RE.match(stripped)
        if heading:
            level = min(len(heading.group(1)), 6)
            text = heading.group(2).strip()
            try:
                para = doc.add_heading("", level=level)
            except (KeyError, ValueError):
                para = doc.add_paragraph()
            _add_runs(para, text)
            n_para += 1
            i += 1
            continue

        # GFM table: a pipe row followed by a separator row
        if "|" in stripped and i + 1 < len(lines) and _TABLE_SEP_RE.match(lines[i + 1]):
            header = _split_row(stripped)
            i += 2
            rows: list[list[str]] = []
            while i < len(lines) and "|" in lines[i] and lines[i].strip():
                rows.append(_split_row(lines[i]))
                i += 1
            width = max([len(header)] + [len(r) for r in rows]) if rows else len(header)
            table = doc.add_table(rows=1, cols=width)
            if table_style:
                table.style = table_style
            for c, value in enumerate(header + [""] * (width - len(header))):
                cell = table.rows[0].cells[c]
                cell.text = ""
                _add_runs(cell.paragraphs[0], value)
                for run in cell.paragraphs[0].runs:
                    run.bold = True
            for row in rows:
                cells = table.add_row().cells
                for c, value in enumerate(row + [""] * (width - len(row))):
                    cells[c].text = ""
                    _add_runs(cells[c].paragraphs[0], value)
            n_table += 1
            continue

        bullet = _BULLET_RE.match(line)
        if bullet:
            para = doc.add_paragraph(style=bullet_style) if bullet_style else doc.add_paragraph()
            _add_runs(para, bullet.group(1))
            n_para += 1
            i += 1
            continue

        number = _NUMBER_RE.match(line)
        if number:
            para = doc.add_paragraph(style=number_style) if number_style else doc.add_paragraph()
            _add_runs(para, number.group(1))
            n_para += 1
            i += 1
            continue

        if stripped.startswith(">"):
            para = doc.add_paragraph(style=quote_style) if quote_style else doc.add_paragraph()
            _add_runs(para, stripped.lstrip("> ").strip())
            n_para += 1
            i += 1
            continue

        # Plain paragraph: join soft-wrapped lines.
        buffer = [stripped]
        i += 1
        while i < len(lines) and lines[i].strip() and not _is_block_start(lines[i]):
            buffer.append(lines[i].strip())
            i += 1
        para = doc.add_paragraph()
        _add_runs(para, " ".join(buffer))
        n_para += 1

    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(target))
    log.info("rendered %s (%d paragraphs, %d tables) from %s", target.name, n_para, n_table, template.name)
    return DocxResult(path=target, paragraphs=n_para, tables=n_table, template=template.name)


def _is_block_start(line: str) -> bool:
    stripped = line.strip()
    return bool(
        _HEADING_RE.match(stripped)
        or _BULLET_RE.match(line)
        or _NUMBER_RE.match(line)
        or _FENCE_RE.match(line)
        or _RULE_RE.match(stripped)
        or stripped.startswith(">")
        or stripped.startswith("|")
    )


def _pt(size: int):
    from docx.shared import Pt

    return Pt(size)
